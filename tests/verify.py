#!/usr/bin/env python3
"""
Build every shipped example, check layout and arithmetic contracts, and prove the
validators still catch the defects they exist for. Exit 0 only if everything passes.

    python3 tests/verify.py            # from the repo root

Needs: node with pptxgenjs, react, react-dom, react-icons, sharp installed globally;
python with reportlab, pypdf, python-pptx.

Three kinds of check:
  1. Positive builds: every example builds with zero WARN lines and exit 0.
  2. Output checks: slide counts, no shape off the 13.333 x 7.5 canvas, one-page PDF,
     dashboard and explorer scripts parse.
  3. Negative tests: a known defect injected into a copy of an example must make the
     validator fire. If a validator goes quiet, the suite fails.

Staleness checks are pinned with --today so results do not drift with the calendar.
"""
import copy, json, os, re, subprocess, sys, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SK = os.path.join(ROOT, "skills")
INTEL = os.path.join(SK, "property-intel-report")
APPR = os.path.join(SK, "property-appraisal")
INS = os.path.join(SK, "property-insurance-brief")
DOSS = os.path.join(INTEL, "examples", "dossier_sample")
TMP = tempfile.mkdtemp(prefix="verify_")

env = dict(os.environ)
try:
    env["NODE_PATH"] = subprocess.check_output(["npm", "root", "-g"], text=True).strip()
except Exception:
    pass

results = []  # (name, ok, detail)


def run(cmd):
    p = subprocess.run(cmd, capture_output=True, text=True, env=env, cwd=ROOT)
    return p.returncode, p.stdout + p.stderr


def record(name, ok, detail=""):
    results.append((name, ok, detail))
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else "\n     " + detail.strip().replace("\n", "\n     ")))


def out(name):
    return os.path.join(TMP, name)


def load(p):
    with open(p) as f:
        return json.load(f)


def dump(obj, name):
    p = out(name)
    with open(p, "w") as f:
        json.dump(obj, f)
    return p


def build_clean(name, cmd):
    code, log = run(cmd)
    warns = [l for l in log.splitlines() if l.startswith("WARN")]
    record(name, code == 0 and not warns, log if (code or warns) else "")
    return code == 0


def expect_warn(name, cmd, pattern):
    """A negative test: the command must print a WARN (or exit non-zero) matching pattern."""
    code, log = run(cmd)
    hit = re.search(pattern, log, re.I)
    record("negative: " + name, bool(hit), f"expected /{pattern}/ but got (exit {code}):\n{log[-600:]}")


# ---------- output checks ----------
def check_pptx(name, path, slides):
    try:
        from pptx import Presentation
    except ImportError:
        record(name + " layout", False, "python-pptx not installed")
        return None
    p = Presentation(path)
    W, H, tol = p.slide_width, p.slide_height, 9144  # 0.01 in
    bad = []
    for i, s in enumerate(p.slides, 1):
        for sh in s.shapes:
            if sh.left is None:
                continue
            if sh.left < -tol or sh.top < -tol or sh.left + sh.width > W + tol or sh.top + sh.height > H + tol:
                bad.append(f"slide {i} '{sh.name}' bottom {(sh.top + sh.height) / 914400:.2f}in right {(sh.left + sh.width) / 914400:.2f}in")
    ok = len(p.slides) == slides and not bad
    record(f"{name} layout ({slides} slides, on canvas)", ok, f"{len(p.slides)} slides; " + "; ".join(bad))
    return p


def slide_text(p, idx):
    return [sh.text_frame.text for sh in p.slides[idx].shapes if sh.has_text_frame]


def check_pdf_one_page(name, path):
    try:
        from pypdf import PdfReader
        n = len(PdfReader(path).pages)
    except Exception as e:
        record(name + " one page", False, str(e))
        return
    record(name + " one page", n == 1, f"{n} pages")


def check_page_script(name, path):
    html = open(path).read()
    scripts = re.findall(r"<script>([\s\S]*?)</script>", html)
    js = out(os.path.basename(path) + ".js")
    with open(js, "w") as f:
        f.write("\n;\n".join(scripts))
    code, log = run(["node", "--check", js])
    record(name + " page script parses", code == 0 and bool(scripts), log)


# ========== 1. positive builds ==========
def node(script, *args):
    return ["node", script] + list(args)


IB, ID = os.path.join(INTEL, "scripts", "build_deck.js"), os.path.join(INTEL, "scripts", "build_dashboard.js")
AB, AE = os.path.join(APPR, "scripts", "build_deck.js"), os.path.join(APPR, "scripts", "build_explorer.js")
NB, NS = os.path.join(INS, "scripts", "build_deck.js"), os.path.join(INS, "scripts", "build_spec_sheet.py")

DF, DP, DA, DI = (os.path.join(DOSS, f) for f in ("property_facts.json", "property_data.json", "appraisal_data.json", "insurance_data.json"))
WILLOW = os.path.join(INTEL, "examples", "willow_bend", "property_data.json")
CREEK = os.path.join(INTEL, "examples", "compare_sample_creek_facts.json")
ASAMPLE = os.path.join(APPR, "examples", "sample_appraisal_data.json")
ISAMPLE = os.path.join(INS, "examples", "illustrative_fairhope", "insurance_data.json")
TODAY = "2026-10-18"

if build_clean("intel deck: willow_bend", node(IB, WILLOW, out("willow.pptx"))):
    check_pptx("intel willow_bend", out("willow.pptx"), 16)
if build_clean("intel deck: dossier_sample", node(IB, DP, out("d_intel.pptx"))):
    check_pptx("intel dossier_sample", out("d_intel.pptx"), 16)
if build_clean("appraisal deck: sample", node(AB, ASAMPLE, out("a_sample.pptx"))):
    check_pptx("appraisal sample", out("a_sample.pptx"), 9)
if build_clean("appraisal deck: dossier_sample", node(AB, DA, out("d_appr.pptx"))):
    check_pptx("appraisal dossier_sample", out("d_appr.pptx"), 9)
if build_clean("comp explorer: sample (strict)", node(AE, ASAMPLE, out("a_sample.html"), "--strict")):
    check_page_script("comp explorer", out("a_sample.html"))
build_clean("comp explorer: dossier_sample (strict)", node(AE, DA, out("d_comps.html"), "--strict"))

ins_default = None
if build_clean("insurance deck: illustrative", node(NB, ISAMPLE, out("i_sample.pptx"))):
    ins_default = check_pptx("insurance illustrative", out("i_sample.pptx"), 11)
if build_clean("insurance deck: dossier_sample", node(NB, DI, out("d_ins.pptx"))):
    check_pptx("insurance dossier_sample", out("d_ins.pptx"), 11)
if build_clean("spec sheet: illustrative", ["python3", NS, ISAMPLE, out("i_sample.pdf")]):
    check_pdf_one_page("spec sheet illustrative", out("i_sample.pdf"))
if build_clean("spec sheet: dossier_sample", ["python3", NS, DI, out("d_spec.pdf")]):
    check_pdf_one_page("spec sheet dossier_sample", out("d_spec.pdf"))

if build_clean("dashboard: dossier_sample, all cross-checks (strict)",
               node(ID, DF, out("d_dash.html"), "--appraisal", DA, "--insurance", DI, "--intel", DP, "--strict", "--today", TODAY)):
    check_page_script("dashboard", out("d_dash.html"))
build_clean("dashboard: legacy 1.1 cache, no repairs or note (strict)", node(ID, CREEK, out("creek_dash.html"), "--strict", "--today", TODAY))

COMPARE = os.path.join(INTEL, "scripts", "build_compare.js")
if os.path.exists(COMPARE):
    if build_clean("compare: dossier_sample vs legacy cache (strict)",
                   node(COMPARE, DF, CREEK, "-o", out("compare.html"), "--strict", "--today", TODAY)):
        check_page_script("compare", out("compare.html"))

# ========== 2. regression: the deductible label ==========
if ins_default is not None:
    t = slide_text(ins_default, 4)
    record("regression: defaulted deductible renders DEFAULT, never YOUR PICK", "DEFAULT" in t and "YOUR PICK" not in t, str(t[:8]))
    picked = load(ISAMPLE)
    for o in picked["wind_deductible"]["options"]:
        o["chosen"] = (o["pct"] == "2%")
    picked["sources"]["caveats"] = [c for c in picked["sources"]["caveats"] if "default" not in c.lower()]
    code, log = run(node(NB, dump(picked, "picked.json"), out("picked.pptx")))
    if code == 0:
        from pptx import Presentation
        t2 = slide_text(Presentation(out("picked.pptx")), 4)
        record("regression: buyer-selected deductible renders YOUR PICK", "YOUR PICK" in t2 and "DEFAULT" not in t2, str(t2[:8]))

# ========== 3. negative tests ==========
d = load(DP); d["comps"]["fair_value_walk"][0]["value"] = "$190 / sqft"
expect_warn("intel comp average that does not reproduce", node(IB, dump(d, "n1.json"), out("n1.pptx")), r"comp average")
d = load(DP); d["monthly_carry"]["scenarios"][0]["total"] += 100
expect_warn("intel carry total that does not sum", node(IB, dump(d, "n2.json"), out("n2.pptx")), r"carry scenario")
d = load(WILLOW); d["red_flags"]["title"] = "Six Things to Resolve"
expect_warn("intel title that hardcodes a count", node(IB, dump(d, "n3.json"), out("n3.pptx")), r"hardcodes a count")

a = load(DA); a["comps"][0]["weight"] = 0.5
expect_warn("appraisal weights not summing to 1", node(AB, dump(a, "n4.json"), out("n4.pptx")), r"weights sum")
a = load(DA); a["value_vs_price"]["verdict_headline"] = "Priced above value"
expect_warn("appraisal verdict mismatch", node(AB, dump(a, "n5.json"), out("n5.pptx")), r"verdict mismatch")
a = load(DA); a["value_conclusion"]["point_value"] += 15000
expect_warn("explorer point that does not reproduce", node(AE, dump(a, "n6.json"), out("n6.html"), "--strict"), r"WARN")

i = load(ISAMPLE)
for o in i["wind_deductible"]["options"]:
    o["chosen"] = True
expect_warn("insurance with three chosen deductibles", node(NB, dump(i, "n7.json"), out("n7.pptx")), r"exactly one wind_deductible")
i = load(ISAMPLE); i["spec_sheet"]["coverage_block"] = i["spec_sheet"]["coverage_block"] * 4
expect_warn("spec sheet left column overflow exits non-zero", ["python3", NS, dump(i, "n8.json"), out("n8.pdf")], r"overflow|exceed|does not fit|too long|footer")

f = load(DF); f["value"]["point"] += 10000
expect_warn("dashboard value drift vs appraisal deck",
            node(ID, dump(f, "n9.json"), out("n9.html"), "--appraisal", DA, "--insurance", DI, "--intel", DP, "--today", TODAY), r"appraisal point_value")
f = load(DF); f.pop("dossier", None)
expect_warn("dashboard full run without the reconciliation note",
            node(ID, dump(f, "n10.json"), out("n10.html"), "--appraisal", DA, "--insurance", DI, "--intel", DP, "--today", TODAY), r"dossier\.note missing")
f = load(DF); f["repairs"]["items"][0]["low"] = f["repairs"]["items"][0]["high"] + 1
expect_warn("dashboard repair item with low above high", node(ID, dump(f, "n11.json"), out("n11.html"), "--today", TODAY), r"low <= high")
f = load(DF); f["market"]["rate_date"] = "2026-09-01"
expect_warn("dashboard stale mortgage rate", node(ID, dump(f, "n12.json"), out("n12.html"), "--today", TODAY), r"days old")
f = load(DF); f["dossier"]["note"]["headline"] = "Worth about $999K."
f["dossier"]["note"]["items"] = [{"label": "x", "body": "no value here"}]
expect_warn("dashboard note that never states the appraisal point", node(ID, dump(f, "n13.json"), out("n13.html"), "--today", TODAY), r"never states the appraisal point")
if os.path.exists(COMPARE):
    f = load(DF); f["address"].update({k: load(CREEK)["address"][k] for k in ("street", "zip")})
    expect_warn("compare with two caches for the same address", node(COMPARE, dump(f, "n14.json"), CREEK, "-o", out("n14.html"), "--today", TODAY), r"same (house|address)")

# ========== hygiene: no real identifiers in shipped examples ==========
ex_text = ""
for base, _, files in os.walk(SK):
    if "examples" in base:
        for fn in files:
            ex_text += open(os.path.join(base, fn)).read()
leaks = re.findall(r"\b\d{3}-\d{3}-\d{4}\b|spokeo|beenverified|truepeoplesearch|BCAR \d{6}|redfin\.com/AL/[^\"]+/home/\d+", ex_text, re.I)
record("hygiene: no phone numbers, aggregators, MLS numbers, or listing URLs in examples", not leaks, ", ".join(sorted(set(leaks))[:5]))

# ========== summary ==========
failed = [r for r in results if not r[1]]
print(f"\n{len(results) - len(failed)}/{len(results)} checks passed" + (f"; {len(failed)} failed" if failed else ""))
sys.exit(1 if failed else 0)
